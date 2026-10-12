"""
بناء ملف Excel لإدارة بنود عقد البوابة الموحدة.

الاستخدام:
    python builder/build_workbook.py data/sample.json output/EduGate_Contract.xlsx

الملف الناتج كله معادلات؛ السكريبت لا يحسب أي رقم بنفسه، فقط يكتب المدخلات والمعادلات والتنسيق.
"""
import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- حجم الجداول
ITEM_ROWS = 40          # عدد صفوف البنود الجاهزة
DELIV_ROWS = 400        # عدد صفوف المخرجات الجاهزة
LOG_ROWS = 1500         # عدد صفوف سجل المراجعات الجاهزة
HOLIDAY_ROWS = 200      # عدد صفوف الأجازات الجاهزة

HEADER_ROW = 4
FIRST = HEADER_ROW + 1  # أول صف بيانات في كل جدول

# ---------------------------------------------------------------- أسماء الشيتات
SH_GUIDE = "دليل الاستخدام"
SH_SETTINGS = "الإعدادات"
SH_ITEMS = "البنود"
SH_DELIV = "المخرجات"
SH_LOG = "سجل المراجعات"
SH_HOLIDAYS = "الأجازات"
SH_LISTS = "القوائم"

# ---------------------------------------------------------------- القيم الثابتة
ACT_SUBMIT = "تسليم"
ACT_REPLY = "رد بملاحظات"
ACT_ACCEPT = "قبول"
ACTIONS = [ACT_SUBMIT, ACT_REPLY, ACT_ACCEPT]

WORK_STATUSES = ["لم يبدأ", "جاري العمل", "معلق"]

ST_ACCEPTED = "تم القبول"
ST_SUSPENDED = "معلق"
ST_RETURNED = "مُعاد بملاحظات"
ST_WAITING = "في انتظار القبول"
ST_LATE = "متأخر"
ST_IN_PROGRESS = "جاري العمل"
ST_NOT_STARTED = "لم يبدأ"

# ---------------------------------------------------------------- التنسيق
FONT = "Arial"
C_INPUT_HEAD = "1F5C6B"     # رأس عمود إدخال
C_CALC_HEAD = "4D4D4D"      # رأس عمود محسوب
C_CALC_FILL = "EFEFEF"      # خلفية الخلايا المحسوبة
C_TITLE = "1F5C6B"
C_TOTAL_FILL = "D9E7EA"

STATUS_COLORS = {
    ST_ACCEPTED: ("C6EFCE", "006100"),
    ST_WAITING: ("DDEBF7", "1F4E78"),
    ST_RETURNED: ("FCE4D6", "843C0C"),
    ST_LATE: ("FFC7CE", "9C0006"),
    ST_SUSPENDED: ("D9D9D9", "404040"),
    ST_IN_PROGRESS: ("FFF2CC", "7F6000"),
}

FMT_MONEY = '#,##0.00;-#,##0.00;"-"'
FMT_INT = '#,##0;-#,##0;"-"'
FMT_PCT = '0.0%;-0.0%;"0%"'
FMT_DATE = 'yyyy-mm-dd'
FMT_DATE_HIDE0 = 'yyyy-mm-dd;;'   # الصفر (= لا يوجد تاريخ) يظهر فارغ

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    return Font(**kw)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def to_date(v):
    return date.fromisoformat(v) if isinstance(v, str) and v else v


# ---------------------------------------------------------------- مساعدات عامة
def rng(sheet, col, first, last):
    return f"'{sheet}'!${col}${first}:${col}${last}"


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def setup_sheet(ws, title, subtitle=None, tab=None):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = font(size=16, bold=True, color=C_TITLE)
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = font(size=10, italic=True, color="595959")
    if tab:
        ws.sheet_properties.tabColor = tab
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_table(ws, columns, n_rows, data_rows):
    """
    columns: قائمة dict لكل عمود:
        key, title, width, kind ('input' | 'calc'),
        formula (نص فيه {r} لرقم الصف) أو None، fmt، wrap، note
    data_rows: قائمة dict بقيم أعمدة الإدخال.
    """
    last = FIRST + n_rows - 1
    for i, col in enumerate(columns, start=1):
        letter = get_column_letter(i)
        col["letter"] = letter
        h = ws.cell(row=HEADER_ROW, column=i, value=col["title"])
        is_input = col["kind"] == "input"
        h.font = font(bold=True, color="FFFFFF")
        h.fill = fill(C_INPUT_HEAD if is_input else C_CALC_HEAD)
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = BORDER
        if col.get("note"):
            h.comment = Comment(col["note"], "النظام", width=320, height=160)
        ws.column_dimensions[letter].width = col.get("width", 14)

    for idx in range(n_rows):
        r = FIRST + idx
        values = data_rows[idx] if idx < len(data_rows) else {}
        for col in columns:
            c = ws[f"{col['letter']}{r}"]
            if col["kind"] == "calc":
                c.value = col["formula"].format(r=r)
                c.fill = fill(C_CALC_FILL)
            else:
                v = values.get(col["key"])
                if col.get("fmt") in (FMT_DATE,):
                    v = to_date(v)
                c.value = v
            c.font = font()
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=col.get("wrap", False),
                                    horizontal=col.get("align", "general"))
            if col.get("fmt"):
                c.number_format = col["fmt"]

    ws.row_dimensions[HEADER_ROW].height = 42
    ws.auto_filter.ref = f"A{HEADER_ROW}:{columns[-1]['letter']}{last}"
    ws.freeze_panes = ws[f"B{FIRST}"]
    return {c["key"]: c["letter"] for c in columns}, last


def add_list_validation(ws, cell_range, source, prompt=None):
    dv = DataValidation(type="list", formula1=source, allow_blank=True)
    dv.error = "اختر قيمة من القائمة"
    dv.errorTitle = "قيمة غير صحيحة"
    dv.showErrorMessage = True
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    for part in cell_range.split():
        dv.add(part)


def add_date_validation(ws, cell_range):
    dv = DataValidation(type="date", operator="between",
                        formula1="DATE(2020,1,1)", formula2="DATE(2040,12,31)", allow_blank=True)
    dv.error = "اكتب تاريخ صحيح (مثال: 2026-01-15)"
    dv.errorTitle = "تاريخ غير صحيح"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    for part in cell_range.split():
        dv.add(part)


def add_number_validation(ws, cell_range, whole=False):
    dv = DataValidation(type="whole" if whole else "decimal", operator="greaterThanOrEqual",
                        formula1="0", allow_blank=True)
    dv.error = "اكتب رقم أكبر من أو يساوي صفر"
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    for part in cell_range.split():
        dv.add(part)


def status_formatting(ws, cell_range, first_cell):
    for status, (bg, fg) in STATUS_COLORS.items():
        ws.conditional_formatting.add(
            cell_range,
            FormulaRule(formula=[f'{first_cell}="{status}"'],
                        fill=fill(bg), font=Font(name=FONT, color=fg, bold=True)))


# ================================================================ الإعدادات
SETTINGS_ROWS = [
    # (key, label, kind, note)
    ("project_name", "اسم المشروع", "input", None),
    ("contract_no", "رقم العقد / كود المشروع", "input", None),
    ("client", "الجهة المالكة", "input", None),
    ("vendor", "المورد / المنفذ", "input", None),
    ("start_date", "تاريخ بداية العقد", "input", None),
    ("duration_months", "مدة العقد (بالأشهر)", "input", None),
    ("end_date", "تاريخ نهاية العقد", "calc", "تاريخ البداية + المدة بالأشهر − يوم"),
    ("vat", "نسبة ضريبة القيمة المضافة", "input", None),
    ("review_days", "مهلة المراجعة (أيام عمل)", "input",
     "عدد أيام العمل المسموح بها للرد على كل تسليم. أي يوم زيادة يُحسب تأخير."),
    ("weekend_code", "رمز الإجازة الأسبوعية", "input",
     "رمز Excel لأيام الإجازة الأسبوعية في دالة NETWORKDAYS.INTL:\n7 = الجمعة والسبت\n1 = السبت والأحد\n16 = الجمعة فقط"),
    ("as_of_override", "تاريخ احتساب ثابت (اختياري)", "input",
     "اتركه فارغاً ليتم الحساب حتى تاريخ اليوم.\nاكتب تاريخاً لو عايز تطلع تقرير \"كما في تاريخ معين\"."),
    ("as_of", "تاريخ الاحتساب الفعلي", "calc", "لو في تاريخ ثابت يُستخدم، وإلا تاريخ اليوم"),
]


def build_settings(wb, settings):
    ws = wb.create_sheet(SH_SETTINGS)
    setup_sheet(ws, "الإعدادات وبيانات العقد",
                "الخلايا البيضاء للإدخال — الرمادية محسوبة تلقائياً", tab="1F5C6B")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 60

    for i, title in enumerate(["البيان", "القيمة", "ملاحظة"]):
        c = ws.cell(row=HEADER_ROW, column=i + 1, value=title)
        c.font = font(bold=True, color="FFFFFF")
        c.fill = fill(C_INPUT_HEAD)
        c.border = BORDER
        c.alignment = Alignment(horizontal="center")

    cells = {}
    for i, (key, label, kind, note) in enumerate(SETTINGS_ROWS):
        r = FIRST + i
        cells[key] = f"'{SH_SETTINGS}'!$B${r}"
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = font(bold=True)
        ws[f"C{r}"] = (note or "").replace("\n", " — ")
        ws[f"C{r}"].font = font(color="595959", size=9)
        ws[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        b = ws[f"B{r}"]
        if kind == "input":
            v = settings.get(key)
            if key in ("start_date", "as_of_override"):
                v = to_date(v)
            b.value = v
        else:
            b.fill = fill(C_CALC_FILL)
        for col in "ABC":
            ws[f"{col}{r}"].border = BORDER
        b.font = font()
        b.alignment = Alignment(horizontal="right")

    row = {key: FIRST + i for i, (key, *_rest) in enumerate(SETTINGS_ROWS)}
    ws[f"B{row['end_date']}"] = f"=IF(ISNUMBER(B{row['start_date']}),EDATE(B{row['start_date']},B{row['duration_months']})-1,\"\")"
    ws[f"B{row['as_of']}"] = f"=IF(ISNUMBER(B{row['as_of_override']}),B{row['as_of_override']},TODAY())"
    for k in ("start_date", "end_date", "as_of_override", "as_of"):
        ws[f"B{row[k]}"].number_format = FMT_DATE
    ws[f"B{row['vat']}"].number_format = "0%"
    add_date_validation(ws, f"B{row['start_date']} B{row['as_of_override']}")
    add_number_validation(ws, f"B{row['duration_months']} B{row['review_days']}", whole=True)

    add_name(wb, "نسبة_الضريبة", cells["vat"])
    add_name(wb, "مهلة_المراجعة", cells["review_days"])
    add_name(wb, "رمز_العطلة", cells["weekend_code"])
    add_name(wb, "تاريخ_الاحتساب", cells["as_of"])
    return ws


# ================================================================ الأجازات
def build_holidays(wb, holidays):
    ws = wb.create_sheet(SH_HOLIDAYS)
    setup_sheet(ws, "الأجازات الرسمية",
                "تُستثنى من حساب أيام المراجعة والتأخير — أضف أي إجازة جديدة في أول صف فاضي", tab="7F7F7F")
    columns = [
        dict(key="date", title="التاريخ", width=16, kind="input", fmt=FMT_DATE),
        dict(key="name", title="المناسبة", width=30, kind="input"),
        dict(key="note", title="ملاحظة", width=24, kind="input"),
    ]
    _, last = write_table(ws, columns, HOLIDAY_ROWS, holidays)
    add_date_validation(ws, f"A{FIRST}:A{last}")
    add_name(wb, "أيام_الأجازات", rng(SH_HOLIDAYS, "A", FIRST, last))
    ws.conditional_formatting.add(
        f"A{FIRST}:A{last}",
        FormulaRule(formula=[f"AND(A{FIRST}<>\"\",COUNTIF($A${FIRST}:$A${last},A{FIRST})>1)"],
                    fill=fill("FFC7CE")))
    return ws


# ================================================================ القوائم
def build_lists(wb):
    ws = wb.create_sheet(SH_LISTS)
    ws.sheet_view.rightToLeft = True
    lists = {"A": ("حالة العمل", WORK_STATUSES, "قائمة_حالة_العمل"),
             "B": ("نوع الحركة", ACTIONS, "قائمة_الحركات")}
    for col, (title, values, name) in lists.items():
        ws[f"{col}1"] = title
        ws[f"{col}1"].font = font(bold=True)
        for i, v in enumerate(values, start=2):
            ws[f"{col}{i}"] = v
        add_name(wb, name, f"'{SH_LISTS}'!${col}$2:${col}${len(values) + 1}")
        ws.column_dimensions[col].width = 20
    ws.sheet_state = "hidden"


# ================================================================ البنود
def build_items(wb, items):
    ws = wb.create_sheet(SH_ITEMS)
    setup_sheet(ws, "بنود العقد",
                "أدخل البنود كما في جدول الكميات — كل الأعمدة الرمادية محسوبة من شيت المخرجات", tab="2E75B6")
    L = DELIV_LETTERS
    deliv = lambda key: f"'{SH_DELIV}'!${L[key]}${FIRST}:${L[key]}${FIRST + DELIV_ROWS - 1}"
    columns = [
        dict(key="no", title="رقم البند", width=10, kind="input", align="center"),
        dict(key="wbs", title="WBS", width=19, kind="input"),
        dict(key="name", title="اسم البند", width=44, kind="input", wrap=True),
        dict(key="qty", title="الكمية التعاقدية", width=11, kind="input", fmt=FMT_INT),
        dict(key="unit_price", title="سعر الوحدة (قبل الضريبة)", width=17, kind="input", fmt=FMT_MONEY),
        dict(key="total", title="الإجمالي قبل الضريبة", width=18, kind="calc", fmt=FMT_MONEY,
             formula='=IF(A{r}="","",D{r}*E{r})', note="الكمية × سعر الوحدة"),
        dict(key="vat", title="الضريبة", width=15, kind="calc", fmt=FMT_MONEY,
             formula='=IF(A{r}="","",F{r}*نسبة_الضريبة)', note="الإجمالي × نسبة الضريبة (من شيت الإعدادات)"),
        dict(key="total_vat", title="الإجمالي شامل الضريبة", width=18, kind="calc", fmt=FMT_MONEY,
             formula='=IF(A{r}="","",F{r}+G{r})'),
        dict(key="registered", title="المخرجات المسجلة", width=11, kind="calc", fmt=FMT_INT,
             formula=f'=IF(A{{r}}="","",COUNTIF({deliv("item")},A{{r}}))',
             note="عدد صفوف المخرجات المسجلة لهذا البند في شيت المخرجات.\nيتلوّن أحمر لو مختلف عن الكمية التعاقدية."),
        dict(key="delivered", title="المسلّمة", width=10, kind="calc", fmt=FMT_INT,
             formula=f'=IF(A{{r}}="","",COUNTIFS({deliv("item")},A{{r}},{deliv("first_submit")},">0"))',
             note="المخرجات اللي اتسلمت مرة واحدة على الأقل"),
        dict(key="accepted", title="المقبولة", width=10, kind="calc", fmt=FMT_INT,
             formula=f'=IF(A{{r}}="","",COUNTIFS({deliv("item")},A{{r}},{deliv("status")},"{ST_ACCEPTED}"))'),
        dict(key="late", title="المتأخرة", width=10, kind="calc", fmt=FMT_INT,
             formula=f'=IF(A{{r}}="","",COUNTIFS({deliv("item")},A{{r}},{deliv("status")},"{ST_LATE}"))'),
        dict(key="progress", title="نسبة الإنجاز", width=11, kind="calc", fmt=FMT_PCT,
             formula='=IF(OR(A{r}="",N(D{r})=0),"",K{r}/D{r})',
             note="عدد المخرجات المقبولة ÷ الكمية التعاقدية"),
        dict(key="accepted_value", title="قيمة المقبول (قبل الضريبة)", width=18, kind="calc", fmt=FMT_MONEY,
             formula='=IF(A{r}="","",K{r}*E{r})'),
        dict(key="remaining_value", title="المتبقي (قبل الضريبة)", width=18, kind="calc", fmt=FMT_MONEY,
             formula='=IF(A{r}="","",F{r}-N{r})'),
    ]
    letters, last = write_table(ws, columns, ITEM_ROWS, items)

    # صف الإجمالي فوق رأس الجدول
    tr = HEADER_ROW - 1
    ws[f"C{tr}"] = "الإجمالي"
    for key in ("qty", "total", "vat", "total_vat", "registered", "delivered", "accepted", "late",
                "accepted_value", "remaining_value"):
        col = letters[key]
        ws[f"{col}{tr}"] = f"=SUM({col}{FIRST}:{col}{last})"
    ws[f"M{tr}"] = f"=IF(SUM(F{FIRST}:F{last})=0,\"\",SUM(N{FIRST}:N{last})/SUM(F{FIRST}:F{last}))"
    ws[f"M{tr}"].comment = Comment(
        "الإنجاز الكلي موزون بالقيمة المالية:\nقيمة كل المخرجات المقبولة ÷ قيمة العقد قبل الضريبة",
        "النظام", width=300, height=110)
    for i, col in enumerate(columns, start=1):
        c = ws.cell(row=tr, column=i)
        c.font = font(bold=True)
        c.fill = fill(C_TOTAL_FILL)
        c.border = BORDER
        if col.get("fmt"):
            c.number_format = col["fmt"]

    add_name(wb, "قائمة_البنود", rng(SH_ITEMS, "A", FIRST, last))
    add_name(wb, "بند_الاسم", rng(SH_ITEMS, "C", FIRST, last))
    add_name(wb, "بند_سعر_الوحدة", rng(SH_ITEMS, "E", FIRST, last))
    add_number_validation(ws, f"D{FIRST}:D{last}", whole=True)
    add_number_validation(ws, f"E{FIRST}:E{last}")

    # تنبيهات: رقم بند مكرر، أو عدد المخرجات المسجلة مختلف عن الكمية
    ws.conditional_formatting.add(
        f"A{FIRST}:A{last}",
        FormulaRule(formula=[f'AND(A{FIRST}<>"",COUNTIF($A${FIRST}:$A${last},A{FIRST})>1)'], fill=fill("FFC7CE")))
    ws.conditional_formatting.add(
        f"I{FIRST}:I{last}",
        FormulaRule(formula=[f'AND(A{FIRST}<>"",I{FIRST}<>D{FIRST})'],
                    fill=fill("FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
    return ws


# ================================================================ المخرجات
DELIV_LETTERS = {}


def deliv_columns():
    log = lambda col: f"'{SH_LOG}'!${col}${FIRST}:${col}${FIRST + LOG_ROWS - 1}"
    LID, LACT, LDATE, LREV, LDELAY = log("A"), log("D"), log("E"), log("L"), log("M")
    blank = '=IF($A{r}="","",'
    return [
        dict(key="id", title="رقم المخرج", width=10, kind="input", align="center",
             note="رقم ثابت ومميز لكل مخرج (مثال: D-001).\nهو اللي بيتربط بيه سجل المراجعات — متغيّروش بعد ما تستخدمه."),
        dict(key="item", title="رقم البند", width=8, kind="input", align="center",
             note="اختر من قائمة بنود العقد"),
        dict(key="item_name", title="اسم البند", width=30, kind="calc", wrap=True,
             formula=blank + 'IFERROR(INDEX(بند_الاسم,MATCH($B{r},قائمة_البنود,0)),"⚠ بند غير موجود"))'),
        dict(key="description", title="وصف المخرج", width=46, kind="input", wrap=True),
        dict(key="planned", title="تاريخ التسليم المخطط", width=13, kind="input", fmt=FMT_DATE,
             note="اختياري. لو اتكتب وعدّى من غير تسليم، الحالة تبقى \"متأخر\" ويتحسب تأخير التسليم."),
        dict(key="work_status", title="حالة العمل (يدوي)", width=12, kind="input",
             note="تُستخدم قبل التسليم فقط: لم يبدأ / جاري العمل.\n\"معلق\" تغلب على أي حالة تانية ما عدا \"تم القبول\"."),
        dict(key="points", title="النقاط المنجزة", width=9, kind="input", fmt='0.##',
             note="اختياري — مثال: عدد الخدمات المنجزة من 25 داخل المخرج"),
        dict(key="value", title="القيمة قبل الضريبة", width=15, kind="calc", fmt=FMT_MONEY,
             formula=blank + 'IFERROR(INDEX(بند_سعر_الوحدة,MATCH($B{r},قائمة_البنود,0)),0))',
             note="سعر الوحدة من شيت البنود"),
        dict(key="value_vat", title="القيمة شامل الضريبة", width=15, kind="calc", fmt=FMT_MONEY,
             formula=blank + 'H{r}*(1+نسبة_الضريبة))'),
        dict(key="first_submit", title="أول تسليم", width=12, kind="calc", fmt=FMT_DATE_HIDE0,
             formula=blank + f'_xlfn.MINIFS({LDATE},{LID},$A{{r}},{LACT},"{ACT_SUBMIT}"))'),
        dict(key="last_submit", title="آخر تسليم", width=12, kind="calc", fmt=FMT_DATE_HIDE0,
             formula=blank + f'_xlfn.MAXIFS({LDATE},{LID},$A{{r}},{LACT},"{ACT_SUBMIT}"))'),
        dict(key="last_reply", title="آخر رد بملاحظات", width=12, kind="calc", fmt=FMT_DATE_HIDE0,
             formula=blank + f'_xlfn.MAXIFS({LDATE},{LID},$A{{r}},{LACT},"{ACT_REPLY}"))'),
        dict(key="accepted", title="تاريخ القبول", width=12, kind="calc", fmt=FMT_DATE_HIDE0,
             formula=blank + f'_xlfn.MAXIFS({LDATE},{LID},$A{{r}},{LACT},"{ACT_ACCEPT}"))'),
        dict(key="cycles", title="عدد مرات التسليم", width=9, kind="calc", fmt=FMT_INT,
             formula=blank + f'COUNTIFS({LID},$A{{r}},{LACT},"{ACT_SUBMIT}"))'),
        dict(key="status", title="الحالة", width=15, kind="calc", align="center",
             formula=(blank
                      + f'IF(M{{r}}>0,"{ST_ACCEPTED}",'
                      + f'IF(F{{r}}="{ST_SUSPENDED}","{ST_SUSPENDED}",'
                      + f'IF(AND(L{{r}}>0,L{{r}}>=K{{r}}),"{ST_RETURNED}",'
                      + f'IF(K{{r}}>0,"{ST_WAITING}",'
                      + f'IF(AND(ISNUMBER(E{{r}}),E{{r}}<تاريخ_الاحتساب),"{ST_LATE}",'
                      + f'IF(F{{r}}="","{ST_NOT_STARTED}",F{{r}})))))))'),
             note="تتحسب تلقائياً بالترتيب ده:\n"
                  "1) في تاريخ قبول ← تم القبول\n"
                  "2) حالة العمل = معلق ← معلق\n"
                  "3) آخر حركة رد بملاحظات ← مُعاد بملاحظات\n"
                  "4) في تسليم ← في انتظار القبول\n"
                  "5) عدّى التاريخ المخطط ← متأخر\n"
                  "6) غير كده ← حالة العمل اليدوية"),
        dict(key="review_days", title="أيام المراجعة", width=10, kind="calc", fmt=FMT_INT,
             formula=blank + f'SUMIFS({LREV},{LID},$A{{r}}))',
             note="مجموع أيام العمل اللي قعدها المخرج عند العميل في كل دورات المراجعة (من سجل المراجعات)"),
        dict(key="review_delay", title="أيام التأخير في القبول", width=11, kind="calc", fmt=FMT_INT,
             formula=blank + f'SUMIFS({LDELAY},{LID},$A{{r}}))',
             note="مجموع الأيام الزيادة عن مهلة المراجعة — كل دورة بتتقاس لوحدها"),
        dict(key="submit_delay", title="أيام تأخير التسليم", width=11, kind="calc", fmt=FMT_INT,
             formula=('=IF(OR($A{r}="",NOT(ISNUMBER(E{r}))),"",'
                      'MAX(0,NETWORKDAYS.INTL(E{r},IF(J{r}>0,J{r},تاريخ_الاحتساب),رمز_العطلة,أيام_الأجازات)-1))'),
             note="أيام العمل من التاريخ المخطط لحد أول تسليم (أو لحد النهارده لو لسه ماتسلمش)"),
        dict(key="notes", title="ملاحظات", width=50, kind="input", wrap=True),
    ]


def build_deliverables(wb, deliverables):
    ws = wb.create_sheet(SH_DELIV)
    setup_sheet(ws, "مخرجات العقد",
                "صف لكل مخرج — التواريخ والحالة والأيام بتتحسب من سجل المراجعات", tab="548235")
    columns = deliv_columns()
    _, last = write_table(ws, columns, DELIV_ROWS, deliverables)

    ws["F2"] = "الرقم التالي المقترح:"
    ws["F2"].font = font(bold=True)
    ws["F2"].alignment = Alignment(horizontal="left")
    ws["G2"] = f'="D-"&TEXT(COUNTA(A{FIRST}:A{last})+1,"000")'
    ws["G2"].font = font(bold=True, color=C_TITLE)

    add_name(wb, "قائمة_المخرجات", rng(SH_DELIV, "A", FIRST, last))
    add_name(wb, "مخرج_البند", rng(SH_DELIV, "B", FIRST, last))
    add_name(wb, "مخرج_الوصف", rng(SH_DELIV, "D", FIRST, last))

    add_list_validation(ws, f"B{FIRST}:B{last}", "=قائمة_البنود")
    add_list_validation(ws, f"F{FIRST}:F{last}", "=قائمة_حالة_العمل")
    add_date_validation(ws, f"E{FIRST}:E{last}")
    add_number_validation(ws, f"G{FIRST}:G{last}")

    status_formatting(ws, f"O{FIRST}:O{last}", f"O{FIRST}")
    ws.conditional_formatting.add(
        f"A{FIRST}:A{last}",
        FormulaRule(formula=[f'AND(A{FIRST}<>"",COUNTIF($A${FIRST}:$A${last},A{FIRST})>1)'], fill=fill("FFC7CE")))
    for col in ("Q", "R"):
        ws.conditional_formatting.add(
            f"{col}{FIRST}:{col}{last}",
            FormulaRule(formula=[f'AND(ISNUMBER({col}{FIRST}),{col}{FIRST}>0)'],
                        font=Font(name=FONT, color="C00000", bold=True)))
    return ws


# ================================================================ سجل المراجعات
def build_log(wb, log):
    ws = wb.create_sheet(SH_LOG)
    setup_sheet(ws, "سجل المراجعات",
                "سجّل كل حركة في صف: تسليم / رد بملاحظات / قبول — الترتيب مش مهم", tab="BF8F00")
    last_row = FIRST + LOG_ROWS - 1
    own = lambda col: f"${col}${FIRST}:${col}${last_row}"
    LID, LACT, LDATE = own("A"), own("D"), own("E")
    is_cycle = '$G{r}=""'
    columns = [
        dict(key="deliverable", title="رقم المخرج", width=10, kind="input", align="center",
             note="اختر من قائمة المخرجات"),
        dict(key="item", title="رقم البند", width=8, kind="calc", align="center",
             formula='=IF($A{r}="","",IFERROR(INDEX(مخرج_البند,MATCH($A{r},قائمة_المخرجات,0)),""))'),
        dict(key="description", title="وصف المخرج", width=40, kind="calc", wrap=True,
             formula='=IF($A{r}="","",IFERROR(INDEX(مخرج_الوصف,MATCH($A{r},قائمة_المخرجات,0)),""))'),
        dict(key="action", title="نوع الحركة", width=13, kind="input", align="center"),
        dict(key="date", title="التاريخ", width=12, kind="input", fmt=FMT_DATE),
        dict(key="notes", title="ملاحظات", width=34, kind="input", wrap=True),
        dict(key="cycle", title="رقم الدورة", width=8, kind="calc", fmt=FMT_INT, align="center",
             formula=(f'=IF(OR($A{{r}}="",$D{{r}}<>"{ACT_SUBMIT}",NOT(ISNUMBER($E{{r}}))),"",'
                      f'COUNTIFS({LID},$A{{r}},{LACT},"{ACT_SUBMIT}",{LDATE},"<"&$E{{r}})+1)'),
             note="ترتيب التسليم ده بين تسليمات نفس المخرج (1 = التسليم الأول)"),
        dict(key="next_answer", title="أول رد/قبول بعده", width=12, kind="calc", fmt=FMT_DATE_HIDE0,
             formula=(f'=IF({is_cycle},"",_xlfn.MINIFS({LDATE},{LID},$A{{r}},'
                      f'{LACT},"<>{ACT_SUBMIT}",{LDATE},">="&$E{{r}}))')),
        dict(key="next_submit", title="التسليم التالي", width=12, kind="calc", fmt=FMT_DATE_HIDE0,
             formula=(f'=IF({is_cycle},"",_xlfn.MINIFS({LDATE},{LID},$A{{r}},'
                      f'{LACT},"{ACT_SUBMIT}",{LDATE},">"&$E{{r}}))')),
        dict(key="cycle_end", title="نهاية الدورة", width=12, kind="calc", fmt=FMT_DATE_HIDE0,
             formula=(f'=IF({is_cycle},"",IF(H{{r}}>0,IF(I{{r}}>0,MIN(H{{r}},I{{r}}),H{{r}}),'
                      'IF(I{r}>0,I{r},تاريخ_الاحتساب)))'),
             note="الدورة بتنتهي عند أقرب حدث بعد التسليم:\n"
                  "رد بملاحظات أو قبول أو تسليم جديد.\n"
                  "لو مفيش أي حدث ← لحد تاريخ الاحتساب (النهارده)."),
        dict(key="cycle_state", title="حالة الدورة", width=9, kind="calc", align="center",
             formula=f'=IF({is_cycle},"",IF(AND(H{{r}}=0,I{{r}}=0),"جارية","مغلقة"))'),
        dict(key="review_days", title="أيام المراجعة", width=9, kind="calc", fmt=FMT_INT,
             formula=(f'=IF({is_cycle},"",MAX(0,NETWORKDAYS.INTL($E{{r}},J{{r}},رمز_العطلة,أيام_الأجازات)-1))'),
             note="أيام العمل من اليوم التالي للتسليم لحد نهاية الدورة.\n"
                  "الإجازة الأسبوعية والأجازات الرسمية مش محسوبة."),
        dict(key="delay_days", title="أيام التأخير", width=9, kind="calc", fmt=FMT_INT,
             formula=f'=IF({is_cycle},"",MAX(0,L{{r}}-مهلة_المراجعة))',
             note="أيام المراجعة − مهلة المراجعة (لو بالسالب = صفر)"),
        dict(key="warning", title="تنبيه", width=22, kind="calc", wrap=True,
             formula=('=IF($A{r}="",IF(OR($D{r}<>"",$E{r}<>""),"رقم المخرج ناقص",""),'
                      'IF(ISNA(MATCH($A{r},قائمة_المخرجات,0)),"رقم مخرج غير موجود",'
                      'IF($D{r}="","نوع الحركة ناقص",'
                      'IF(NOT(ISNUMBER($E{r})),"التاريخ ناقص",'
                      'IF($E{r}>تاريخ_الاحتساب,"تاريخ بعد تاريخ الاحتساب",'
                      f'IF(AND($D{{r}}<>"{ACT_SUBMIT}",COUNTIFS({LID},$A{{r}},{LACT},"{ACT_SUBMIT}",{LDATE},"<="&$E{{r}})=0),'
                      '"لا يوجد تسليم قبل الحركة دي",""))))))')),
    ]
    _, last = write_table(ws, columns, LOG_ROWS, log)
    assert last == last_row

    add_list_validation(ws, f"A{FIRST}:A{last}", "=قائمة_المخرجات")
    add_list_validation(ws, f"D{FIRST}:D{last}", "=قائمة_الحركات")
    add_date_validation(ws, f"E{FIRST}:E{last}")

    add_name(wb, "سجل_المخرج", rng(SH_LOG, "A", FIRST, last))
    ws.conditional_formatting.add(
        f"M{FIRST}:M{last}",
        FormulaRule(formula=[f'AND(ISNUMBER(M{FIRST}),M{FIRST}>0)'], font=Font(name=FONT, color="C00000", bold=True)))
    ws.conditional_formatting.add(
        f"N{FIRST}:N{last}",
        FormulaRule(formula=[f'N{FIRST}<>""'], fill=fill("FFC7CE"), font=Font(name=FONT, color="9C0006", bold=True)))
    ws.conditional_formatting.add(
        f"K{FIRST}:K{last}",
        FormulaRule(formula=[f'K{FIRST}="جارية"'], fill=fill("FFF2CC")))
    return ws


# ================================================================ دليل الاستخدام
GUIDE = [
    ("h", "طريقة العمل"),
    ("p", "1) شيت «الإعدادات»: بيانات العقد، نسبة الضريبة، مهلة المراجعة، والإجازة الأسبوعية."),
    ("p", "2) شيت «البنود»: كل بند تعاقدي في صف — الرقم، WBS، الاسم، الكمية، سعر الوحدة."),
    ("p", "3) شيت «المخرجات»: صف لكل مخرج تحت البند بتاعه، برقم ثابت (D-001، D-002، ...)."),
    ("p", "4) شيت «سجل المراجعات»: كل ما يحصل تسليم أو رد بملاحظات أو قبول، ضيف صف جديد بالتاريخ."),
    ("p", "5) شيت «الأجازات»: أي إجازة رسمية جديدة تضاف هنا عشان تتشال من حساب الأيام."),
    ("h", "ألوان الخلايا"),
    ("legend_input", "رأس عمود أزرق = عمود إدخال، اكتب فيه."),
    ("legend_calc", "رأس عمود رمادي وخلفية رمادية = معادلة، متكتبش فيه."),
    ("p", "حط الماوس على رأس أي عمود عليه مثلث أحمر صغير عشان تشوف شرح المعادلة."),
    ("h", "حساب الحالة"),
    ("p", "الحالة مش بتتكتب يدوي — بتتحسب من سجل المراجعات:"),
    ("p", "• في حركة «قبول» ← تم القبول"),
    ("p", "• حالة العمل اليدوية = معلق ← معلق"),
    ("p", "• آخر حركة «رد بملاحظات» ← مُعاد بملاحظات"),
    ("p", "• آخر حركة «تسليم» ← في انتظار القبول"),
    ("p", "• عدّى تاريخ التسليم المخطط من غير تسليم ← متأخر"),
    ("p", "• غير كده ← حالة العمل اليدوية (لم يبدأ / جاري العمل)"),
    ("h", "حساب أيام المراجعة والتأخير"),
    ("p", "• كل «تسليم» بيفتح دورة مراجعة، والدورة بتقفل عند أول رد أو قبول أو تسليم جديد بعده."),
    ("p", "• الأيام = أيام عمل من اليوم التالي للتسليم، من غير الإجازة الأسبوعية والأجازات الرسمية."),
    ("p", "• التأخير = الزيادة عن مهلة المراجعة في كل دورة لوحدها (توفير دورة مش بيعوّض تأخير دورة تانية)."),
    ("p", "• لو الدورة لسه مفتوحة، الحساب بيكمّل لحد تاريخ الاحتساب (النهارده أو التاريخ الثابت في الإعدادات)."),
    ("h", "مثال لتسجيل مخرج في سجل المراجعات"),
    ("p", "D-005 | تسليم | 2026-03-15"),
    ("p", "D-005 | رد بملاحظات | 2026-03-29"),
    ("p", "D-005 | تسليم | 2026-04-05      ← إعادة التسليم = تسليم جديد"),
    ("p", "D-005 | قبول | 2026-04-12"),
]


def build_guide(wb):
    ws = wb.active
    ws.title = SH_GUIDE
    setup_sheet(ws, "نظام إدارة بنود عقد البوابة الموحدة", "دليل الاستخدام", tab="C00000")
    ws.column_dimensions["A"].width = 110
    r = 4
    for kind, text in GUIDE:
        if kind == "h":
            r += 1  # سطر فاضي قبل كل عنوان
        c = ws[f"A{r}"]
        c.value = text
        if kind == "h":
            c.font = font(size=12, bold=True, color=C_TITLE)
        elif kind == "legend_input":
            c.font = font(bold=True, color="FFFFFF")
            c.fill = fill(C_INPUT_HEAD)
        elif kind == "legend_calc":
            c.font = font(bold=True, color="FFFFFF")
            c.fill = fill(C_CALC_HEAD)
        else:
            c.font = font(size=11)
        c.alignment = Alignment(horizontal="right")
        r += 1


# ================================================================ main
def build(data, out_path):
    wb = Workbook()
    build_guide(wb)
    build_settings(wb, data.get("settings", {}))

    global DELIV_LETTERS
    DELIV_LETTERS = {c["key"]: get_column_letter(i) for i, c in enumerate(deliv_columns(), start=1)}

    build_items(wb, data.get("items", []))
    build_deliverables(wb, data.get("deliverables", []))
    build_log(wb, data.get("log", []))
    build_holidays(wb, data.get("holidays", []))
    build_lists(wb)

    for n, rows in (("items", ITEM_ROWS), ("deliverables", DELIV_ROWS),
                    ("log", LOG_ROWS), ("holidays", HOLIDAY_ROWS)):
        if len(data.get(n, [])) > rows:
            raise SystemExit(f"عدد صفوف {n} أكبر من الحد المجهز ({rows})")

    wb.active = 0
    wb.save(out_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("الاستخدام: python builder/build_workbook.py <data.json> <output.xlsx>")
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2])
    print("تم إنشاء", sys.argv[2])
