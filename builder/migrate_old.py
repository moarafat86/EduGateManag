"""
نقل البيانات من ملف المتابعة القديم (Edu_Gate_Management_Sheet.xlsx) إلى ملف JSON
بالشكل اللي يفهمه build_workbook.py.

الاستخدام:
    python builder/migrate_old.py <old.xlsx> data/real.json

ملاحظة: الملف الناتج فيه بيانات العقد الحقيقية — متعملوش commit (موجود في .gitignore).
"""
import json
import sys
from datetime import datetime

from openpyxl import load_workbook

MAIN_SHEET = "مخرجات المشروع"
HOLIDAY_SHEET = "الإعدادات والأجازات"

# أعمدة الملف القديم
COL_ITEM_NO, COL_WBS, COL_ITEM_NAME, COL_QTY, COL_PRICE = "B", "C", "D", "E", "F"
COL_DESC, COL_POINTS, COL_STATUS, COL_NOTES = "I", "J", "K", "T"
LEGACY_BILLING = {"L": "certificate", "M": "invoiced", "N": "ref_no", "P": "paid",
                  "Q": "acceptance_form", "R": "invoice_no", "S": "claim_no"}
# ترتيب أعمدة التواريخ: تسليم ← رد ← إعادة تسليم ← ... ← قبول
DATE_COLUMNS = [("U", "تسليم"), ("V", "رد بملاحظات"), ("W", "تسليم"), ("X", "رد بملاحظات"),
                ("Y", "تسليم"), ("Z", "رد بملاحظات"), ("AA", "تسليم"), ("AB", "قبول")]

# الحالة القديمة ← حالة العمل اليدوية في التصميم الجديد
# (تم / في انتظار القبول بتتحسب من سجل المراجعات، و"متأخر" بتتحسب من تاريخ التسليم المخطط)
WORK_STATUS_MAP = {"معلق": "معلق", "جاري العمل": "جاري العمل", "متأخر": "جاري العمل",
                   "تم": None, "في انتظار القبول": None}


def clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def iso(v):
    return v.date().isoformat() if isinstance(v, datetime) else None


def migrate(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[MAIN_SHEET]
    table_last_row = int(ws.tables["Table3"].ref.split(":")[1].lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))

    items, deliverables, log, warnings = [], [], [], []
    current_item = None
    for r in range(2, table_last_row + 1):
        cell = lambda col: clean(ws[f"{col}{r}"].value)
        if cell(COL_ITEM_NO) is not None:
            current_item = cell(COL_ITEM_NO)
            items.append({"no": current_item, "wbs": cell(COL_WBS), "name": cell(COL_ITEM_NAME),
                          "qty": cell(COL_QTY), "unit_price": cell(COL_PRICE)})
        if current_item is None or cell(COL_DESC) is None:
            continue

        did = f"D-{len(deliverables) + 1:03d}"
        old_status = cell(COL_STATUS)
        deliverables.append({
            "id": did, "item": current_item, "description": cell(COL_DESC),
            "work_status": WORK_STATUS_MAP.get(old_status), "points": cell(COL_POINTS),
            "notes": cell(COL_NOTES), "old_row": r, "old_status": old_status,
            "legacy_billing": {k: cell(c) for c, k in LEGACY_BILLING.items() if cell(c) is not None},
        })

        dates = [(col, act, iso(ws[f"{col}{r}"].value)) for col, act in DATE_COLUMNS]
        for col, act, d in dates:
            if d:
                log.append({"deliverable": did, "action": act, "date": d})
        accepted = any(act == "قبول" and d for _, act, d in dates)
        if old_status == "تم" and not accepted:
            warnings.append(f"صف {r} ({did}): الحالة القديمة «تم» لكن مفيش تاريخ قبول")
        if old_status != "تم" and accepted:
            warnings.append(f"صف {r} ({did}): في تاريخ قبول لكن الحالة القديمة «{old_status}»")

    hs = wb[HOLIDAY_SHEET]
    holidays = [{"date": iso(hs[f"A{r}"].value), "name": clean(hs[f"B{r}"].value),
                 "note": clean(hs[f"C{r}"].value)}
                for r in range(7, hs.max_row + 1) if iso(hs[f"A{r}"].value)]

    settings = {
        "project_name": "البوابة الموحدة للخدمات",
        "contract_no": "CX-2023-0018",
        "client": None,
        "vendor": None,
        "start_date": "2024-04-01",
        "duration_months": 30,
        "vat": 0.15,
        "review_days": clean(hs["B2"].value) or 10,
        "weekend_code": 7,
        "as_of_override": None,
    }
    return {"settings": settings, "items": items, "deliverables": deliverables, "log": log,
            "holidays": holidays}, warnings


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("الاستخدام: python builder/migrate_old.py <old.xlsx> <out.json>")
    data, warnings = migrate(sys.argv[1])
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"بنود: {len(data['items'])} | مخرجات: {len(data['deliverables'])} | "
          f"حركات: {len(data['log'])} | أجازات: {len(data['holidays'])}")
    for w in warnings:
        print("تنبيه:", w)
